from collections import defaultdict
from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO
from pathlib import Path
import textwrap
import unicodedata

from django.contrib import messages
from django.core.exceptions import ValidationError
from django.core.paginator import Paginator
from django.db import DatabaseError
from django.db.models import Q
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.http import require_POST
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from PIL import Image, ImageDraw, ImageFont

from config.permissions import group_required
from empresas.decorators import empresa_required
from documentos.excel import ExcelColumn, ExcelReportBuilder
from documentos.formatting import format_date_br, format_money_br
from documentos.pdf import PdfDocument, PdfTableColumn
from empresas.documentos import draw_empresa_footer

from .forms import (
    CentroCustoForm,
    ContaPagarBaixaForm,
    ContaPagarForm,
    ContaReceberBaixaForm,
    ContaReceberForm,
    FinanceiroFiltroForm,
    FornecedorForm,
    ImportarCredoresSiengeForm,
    ItemContaPagarOrdemCompraFormSet,
    PrevisaoFinanceiraForm,
)
from .importadores import decodificar_csv_upload, importar_contas_pagar_credores_csv, importar_contas_pagas_credores_csv
from .models import CentroCusto, ContaPagar, ContaReceber, Fornecedor, PrevisaoFinanceira
from .services import baixar_conta_pagar as baixar_conta_pagar_service
from .services import baixar_conta_receber as baixar_conta_receber_service


financeiro_required = group_required('Financeiro', 'Diretoria')
RELATORIO_FINANCEIRO_COLUNAS = dict(FinanceiroFiltroForm.COLUNAS_CHOICES)


def _status_visual(conta, tipo):
    if conta.status == 'cancelado':
        return 'Cancelado'
    if tipo == 'receber' and conta.status == ContaReceber.STATUS_RECEBIDO:
        return 'Recebido'
    if tipo == 'pagar' and conta.status == ContaPagar.STATUS_PAGO:
        return 'Pago'
    if conta.data_vencimento < timezone.localdate():
        return 'Atrasado'
    return 'Em aberto'


def _base_receber(empresa):
    return ContaReceber.objects.filter(empresa=empresa).select_related('obra', 'centro_custo', 'nota_fiscal')


def _base_pagar(empresa):
    return ContaPagar.objects.filter(empresa=empresa).select_related('obra', 'centro_custo', 'despesa_obra')


def _base_previsoes(empresa):
    return PrevisaoFinanceira.objects.filter(empresa=empresa).select_related('obra', 'centro_custo')


def _paginar(request, queryset, per_page=50):
    paginator = Paginator(queryset, per_page)
    return paginator.get_page(request.GET.get('page'))


def _formset_tem_itens_oc(formset):
    return any(
        form.cleaned_data.get('item_ordem_compra') and not form.cleaned_data.get('DELETE')
        for form in formset.forms
        if hasattr(form, 'cleaned_data')
    )


def _filtrar_contas(request, data=None):
    form = FinanceiroFiltroForm(data if data is not None else request.GET or None, empresa=request.empresa)
    receber = _base_receber(request.empresa)
    pagar = _base_pagar(request.empresa)
    previsoes = _base_previsoes(request.empresa)
    tipo = status = ''

    if form.is_valid():
        tipo = form.cleaned_data['tipo']
        status = form.cleaned_data['status']
        data_inicial = form.cleaned_data['data_inicial']
        data_final = form.cleaned_data['data_final']
        obra = form.cleaned_data['obra']
        centro_custo = form.cleaned_data['centro_custo']
        busca = form.cleaned_data['busca']

        if data_inicial:
            receber = receber.filter(data_vencimento__gte=data_inicial)
            pagar = pagar.filter(data_vencimento__gte=data_inicial)
            previsoes = previsoes.filter(data_prevista__gte=data_inicial)
        if data_final:
            receber = receber.filter(data_vencimento__lte=data_final)
            pagar = pagar.filter(data_vencimento__lte=data_final)
            previsoes = previsoes.filter(data_prevista__lte=data_final)
        if obra:
            receber = receber.filter(obra__nome_obra__icontains=obra)
            pagar = pagar.filter(obra__nome_obra__icontains=obra)
            previsoes = previsoes.filter(obra__nome_obra__icontains=obra)
        if centro_custo:
            receber = receber.filter(centro_custo=centro_custo)
            pagar = pagar.filter(centro_custo=centro_custo)
            previsoes = previsoes.filter(centro_custo=centro_custo)
        if busca:
            receber = receber.filter(Q(cliente__icontains=busca) | Q(descricao__icontains=busca) | Q(numero_nf__icontains=busca))
            pagar = pagar.filter(Q(fornecedor__icontains=busca) | Q(descricao__icontains=busca))
            previsoes = previsoes.filter(Q(pessoa__icontains=busca) | Q(descricao__icontains=busca))
        if status == 'aberto':
            receber = receber.filter(status=ContaReceber.STATUS_ABERTO)
            pagar = pagar.filter(status=ContaPagar.STATUS_ABERTO)
            previsoes = previsoes.filter(status=PrevisaoFinanceira.STATUS_ATIVA)
        elif status == 'baixado':
            receber = receber.filter(status=ContaReceber.STATUS_RECEBIDO)
            pagar = pagar.filter(status=ContaPagar.STATUS_PAGO)
            previsoes = previsoes.filter(status=PrevisaoFinanceira.STATUS_REALIZADA)
        elif status == 'cancelado':
            receber = receber.filter(status=ContaReceber.STATUS_CANCELADO)
            pagar = pagar.filter(status=ContaPagar.STATUS_CANCELADO)
            previsoes = previsoes.filter(status=PrevisaoFinanceira.STATUS_CANCELADA)
        elif status == 'atrasado':
            hoje = timezone.localdate()
            receber = receber.filter(status=ContaReceber.STATUS_ABERTO, data_vencimento__lt=hoje)
            pagar = pagar.filter(status=ContaPagar.STATUS_ABERTO, data_vencimento__lt=hoje)
            previsoes = previsoes.filter(status=PrevisaoFinanceira.STATUS_ATIVA, data_prevista__lt=hoje)

    if tipo == 'receber':
        pagar = _base_pagar(request.empresa).none()
        previsoes = _base_previsoes(request.empresa).none()
    elif tipo == 'pagar':
        receber = _base_receber(request.empresa).none()
        previsoes = _base_previsoes(request.empresa).none()
    elif tipo == 'previsao':
        receber = _base_receber(request.empresa).none()
        pagar = _base_pagar(request.empresa).none()

    return form, receber, pagar, previsoes


def _eventos_fluxo(receber, pagar, previsoes=None):
    eventos = []
    for conta in receber:
        eventos.append(
            {
                'data': conta.data_recebimento or conta.data_vencimento,
                'tipo': 'Receber',
                'descricao': conta.descricao,
                'pessoa': conta.cliente,
                'obra': conta.obra,
                'centro_custo': conta.centro_custo,
                'status': _status_visual(conta, 'receber'),
                'valor_abs': conta.valor_liquido,
                'valor': conta.valor_liquido,
            }
        )
    for conta in pagar:
        eventos.append(
            {
                'data': conta.data_pagamento or conta.data_vencimento,
                'tipo': 'Pagar',
                'descricao': conta.descricao,
                'pessoa': conta.fornecedor,
                'obra': conta.obra,
                'centro_custo': conta.centro_custo,
                'status': _status_visual(conta, 'pagar'),
                'valor_abs': conta.valor_pago_efetivo if conta.status == ContaPagar.STATUS_PAGO else conta.valor,
                'valor': -conta.valor_pago_efetivo if conta.status == ContaPagar.STATUS_PAGO else -conta.valor,
            }
        )
    for previsao in previsoes or []:
        tipo_label = 'Previsao entrada' if previsao.tipo == PrevisaoFinanceira.TIPO_RECEBER else 'Previsao saida'
        eventos.append(
            {
                'data': previsao.data_prevista,
                'tipo': tipo_label,
                'descricao': previsao.descricao,
                'pessoa': previsao.pessoa,
                'obra': previsao.obra,
                'centro_custo': previsao.centro_custo,
                'status': previsao.get_status_display(),
                'valor_abs': previsao.valor,
                'valor': previsao.valor_fluxo,
            }
        )
    return sorted(eventos, key=lambda item: (item['data'], item['tipo'], item['descricao']))


def _ordenar_eventos(eventos, ordenacao):
    ordenacao = ordenacao or 'data_asc'
    key_map = {
        'data_asc': lambda item: (item['data'], item['pessoa'], item['descricao']),
        'data_desc': lambda item: (item['data'], item['pessoa'], item['descricao']),
        'fornecedor': lambda item: ((item['pessoa'] or '').lower(), item['data'], item['descricao']),
        'centro_custo': lambda item: (str(item['centro_custo'] or '').lower(), item['data'], item['pessoa']),
        'obra': lambda item: (str(item['obra'] or '').lower(), item['data'], item['pessoa']),
        'valor_desc': lambda item: (item['valor_abs'], item['data']),
        'valor_asc': lambda item: (item['valor_abs'], item['data']),
    }
    reverse = ordenacao in {'data_desc', 'valor_desc'}
    return sorted(eventos, key=key_map.get(ordenacao, key_map['data_asc']), reverse=reverse)


def _grupo_evento(evento, agrupamento):
    if agrupamento == 'centro_custo':
        return str(evento['centro_custo'] or 'Sem centro de custo')
    if agrupamento == 'fornecedor':
        return evento['pessoa'] or 'Sem fornecedor/cliente'
    if agrupamento == 'obra':
        return str(evento['obra'] or 'Sem obra')
    if agrupamento == 'status':
        return evento['status'] or 'Sem status'
    if agrupamento == 'tipo':
        return evento['tipo'] or 'Sem tipo'
    return 'Lancamentos'


def _agrupar_eventos(eventos, agrupamento):
    if not agrupamento:
        return [{'titulo': 'Lancamentos', 'eventos': eventos, 'total': sum((e['valor'] for e in eventos), Decimal('0'))}]
    grupos = []
    indices = {}
    for evento in eventos:
        titulo = _grupo_evento(evento, agrupamento)
        if titulo not in indices:
            indices[titulo] = len(grupos)
            grupos.append({'titulo': titulo, 'eventos': [], 'total': Decimal('0')})
        grupo = grupos[indices[titulo]]
        grupo['eventos'].append(evento)
        grupo['total'] += evento['valor']
    return grupos


def _resumo(receber, pagar, previsoes=None):
    hoje = timezone.localdate()
    receber = list(receber)
    pagar = list(pagar)
    previsoes = list(previsoes or [])
    total_receber_aberto = sum((c.valor_liquido for c in receber if c.status == ContaReceber.STATUS_ABERTO), Decimal('0'))
    total_recebido = sum((c.valor_liquido for c in receber if c.status == ContaReceber.STATUS_RECEBIDO), Decimal('0'))
    total_pagar_aberto = sum((c.valor for c in pagar if c.status == ContaPagar.STATUS_ABERTO), Decimal('0'))
    total_pago = sum((c.valor_pago_efetivo for c in pagar if c.status == ContaPagar.STATUS_PAGO), Decimal('0'))
    total_previsao_receber = sum(
        (
            p.valor
            for p in previsoes
            if p.status == PrevisaoFinanceira.STATUS_ATIVA and p.tipo == PrevisaoFinanceira.TIPO_RECEBER
        ),
        Decimal('0'),
    )
    total_previsao_pagar = sum(
        (
            p.valor
            for p in previsoes
            if p.status == PrevisaoFinanceira.STATUS_ATIVA and p.tipo == PrevisaoFinanceira.TIPO_PAGAR
        ),
        Decimal('0'),
    )
    atrasado_receber = sum(
        (c.valor_liquido for c in receber if c.status == ContaReceber.STATUS_ABERTO and c.data_vencimento < hoje),
        Decimal('0'),
    )
    atrasado_pagar = sum(
        (c.valor for c in pagar if c.status == ContaPagar.STATUS_ABERTO and c.data_vencimento < hoje),
        Decimal('0'),
    )
    return {
        'total_receber_aberto': total_receber_aberto,
        'total_recebido': total_recebido,
        'total_pagar_aberto': total_pagar_aberto,
        'total_pago': total_pago,
        'total_previsao_receber': total_previsao_receber,
        'total_previsao_pagar': total_previsao_pagar,
        'saldo_previsto': total_receber_aberto - total_pagar_aberto,
        'saldo_com_previsoes': total_receber_aberto - total_pagar_aberto + total_previsao_receber - total_previsao_pagar,
        'saldo_realizado': total_recebido - total_pago,
        'atrasado_receber': atrasado_receber,
        'atrasado_pagar': atrasado_pagar,
    }


def _grafico_fluxo(eventos):
    semanas = defaultdict(lambda: {'receber': Decimal('0'), 'pagar': Decimal('0')})
    for evento in eventos:
        inicio_semana = evento['data'] - timedelta(days=evento['data'].weekday())
        fim_semana = inicio_semana + timedelta(days=6)
        chave = (inicio_semana, fim_semana)
        if evento['tipo'] in {'Receber', 'Previsao entrada'}:
            semanas[chave]['receber'] += evento['valor']
        else:
            semanas[chave]['pagar'] += abs(evento['valor'])
    labels = sorted(semanas.keys())
    return {
        'labels': [f'{inicio.strftime("%d/%m")} a {fim.strftime("%d/%m")}' for inicio, fim in labels],
        'receber': [float(semanas[label]['receber']) for label in labels],
        'pagar': [float(semanas[label]['pagar']) for label in labels],
    }


def _format_currency_br(value):
    value = value or Decimal('0')
    formatted = f'{value:,.2f}'.replace(',', 'X').replace('.', ',').replace('X', '.')
    return f'R$ {formatted}'


def _pdf_font(size, bold=False):
    font_dir = Path('C:/Windows/Fonts')
    candidates = ['arialbd.ttf' if bold else 'arial.ttf', 'calibrib.ttf' if bold else 'calibri.ttf']
    for name in candidates:
        path = font_dir / name
        if path.exists():
            return ImageFont.truetype(str(path), size)
    return ImageFont.load_default()


def _clean_pdf_text(value):
    return unicodedata.normalize('NFKD', str(value or '-')).encode('ascii', 'ignore').decode('ascii')


def _draw_report_cell(
    draw,
    value,
    x,
    y,
    w,
    h,
    font,
    fill=(17, 24, 39),
    bg=None,
    border=(156, 163, 175),
    align='left',
    width=1,
):
    if bg:
        draw.rectangle((x, y, x + w, y + h), fill=bg, outline=border, width=width)
    else:
        draw.rectangle((x, y, x + w, y + h), outline=border, width=width)
    text = _clean_pdf_text(value)
    avg_char_width = max(font.getlength('ABCDEFGHIJKLMNOPQRSTUVWXYZabcdefghijklmnopqrstuvwxyz') / 52, 1)
    max_chars = max(int((w - 8) / avg_char_width), 4)
    lines = textwrap.wrap(text, width=max_chars) or ['']
    line_height = font.getbbox('Ag')[3] - font.getbbox('Ag')[1] + 4
    visible_lines = lines[: max(int((h - 6) / line_height), 1)]
    text_h = line_height * len(visible_lines)
    y_text = y + max((h - text_h) // 2, 3)
    for line in visible_lines:
        if align == 'right':
            x_text = x + w - font.getlength(line) - 5
        elif align == 'center':
            x_text = x + (w - font.getlength(line)) / 2
        else:
            x_text = x + 5
        draw.text((x_text, y_text), line, font=font, fill=fill)
        y_text += line_height


def _relatorio_financeiro_linha_display(linha, coluna):
    if coluna == 'data':
        return linha['data'].strftime('%d/%m/%Y') if linha.get('data') else '-'
    if coluna == 'valor':
        return format_money_br(linha.get('valor'))
    return str(linha.get(coluna) or '-')


def _relatorio_financeiro_excel_value(linha, coluna):
    if coluna == 'data':
        return linha.get('data')
    if coluna == 'valor':
        return linha.get('valor') or Decimal('0')
    return linha.get(coluna) or '-'


def _xlsx_relatorio_financeiro(eventos, colunas, empresa):
    excel_columns = []
    for coluna in colunas:
        width = 16
        align = 'center'
        number_format = None
        wrap = False
        if coluna in {'descricao', 'pessoa', 'obra', 'centro_custo'}:
            width = 34
            align = 'left'
            wrap = True
        elif coluna == 'valor':
            width = 18
            align = 'right'
            number_format = ExcelReportBuilder.MONEY_FORMAT
        elif coluna == 'data':
            width = 15
            number_format = ExcelReportBuilder.DATE_FORMAT
        excel_columns.append(
            ExcelColumn(
                key=coluna,
                label=RELATORIO_FINANCEIRO_COLUNAS[coluna],
                width=width,
                align=align,
                number_format=number_format,
                wrap=wrap,
            )
        )
    rows = [
        {coluna: _relatorio_financeiro_excel_value(evento, coluna) for coluna in colunas}
        for evento in eventos
    ]
    builder = ExcelReportBuilder(empresa=empresa, title='Relatorio gerencial financeiro', sheet_name='Financeiro')
    builder.add_header(emitted_on=date.today())
    builder.add_table(excel_columns, rows)
    return builder.build()


def _financial_report_pdf(eventos, colunas, resumo, empresa):
    colunas = colunas or FinanceiroFiltroForm.COLUNAS_PADRAO
    weights = {
        'tipo': 1.1,
        'data': 0.9,
        'descricao': 2.2,
        'pessoa': 1.8,
        'obra': 1.8,
        'centro_custo': 1.6,
        'status': 1.0,
        'valor': 1.1,
    }
    pdf_columns = []
    for coluna in colunas:
        align = 'center'
        if coluna in {'descricao', 'pessoa', 'obra', 'centro_custo'}:
            align = 'left'
        elif coluna == 'valor':
            align = 'right'
        pdf_columns.append(
            PdfTableColumn(
                key=coluna,
                label=RELATORIO_FINANCEIRO_COLUNAS[coluna],
                weight=weights.get(coluna, 1),
                align=align,
            )
        )
    rows = [
        {coluna: _relatorio_financeiro_linha_display(evento, coluna) for coluna in colunas}
        for evento in eventos
    ]
    doc = PdfDocument(
        empresa=empresa,
        title='Relatorio gerencial financeiro',
        subtitle=f'{len(eventos)} registro(s)',
        orientation='landscape',
        filename='relatorio_financeiro.pdf',
    )
    doc.add_title(emitted_on=date.today())
    doc.add_info_grid(
        [
            ('A receber aberto', format_money_br(resumo['total_receber_aberto'])),
            ('A pagar aberto', format_money_br(resumo['total_pagar_aberto'])),
            ('Saldo previsto', format_money_br(resumo['saldo_com_previsoes'])),
            ('Saldo realizado', format_money_br(resumo['saldo_realizado'])),
        ],
        columns=4,
    )
    doc.add_table(pdf_columns, rows, row_height=44)
    return doc.build()


@financeiro_required
@empresa_required
def financeiro_home(request):
    receber = _base_receber(request.empresa)
    pagar = _base_pagar(request.empresa)
    previsoes = _base_previsoes(request.empresa)
    eventos = _eventos_fluxo(receber, pagar, previsoes)
    hoje = timezone.localdate()
    limite_fluxo = hoje + timedelta(days=90)
    eventos_fluxo = [evento for evento in eventos if hoje <= evento['data'] <= limite_fluxo]
    contexto = {
        **_resumo(receber, pagar, previsoes),
        'ultimos_eventos': eventos[-10:],
        'grafico_fluxo': _grafico_fluxo(eventos_fluxo),
        'fluxo_periodo': f'{hoje.strftime("%d/%m/%Y")} a {limite_fluxo.strftime("%d/%m/%Y")}',
    }
    return render(request, 'financeiro/home.html', contexto)


@financeiro_required
@empresa_required
def lista_contas_receber(request):
    contas = _base_receber(request.empresa)
    page_obj = _paginar(request, contas)
    return render(request, 'financeiro/lista_contas_receber.html', {'contas': page_obj, 'page_obj': page_obj})


@financeiro_required
@empresa_required
def lista_contas_pagar(request):
    contas = _base_pagar(request.empresa).filter(status=ContaPagar.STATUS_ABERTO)
    page_obj = _paginar(request, contas)
    return render(
        request,
        'financeiro/lista_contas_pagar.html',
        {
            'contas': page_obj,
            'page_obj': page_obj,
            'titulo': 'Contas a Pagar',
            'descricao': 'Despesas em aberto para baixa ou cancelamento em massa.',
            'mostrar_acoes_massa': True,
        },
    )


@financeiro_required
@empresa_required
def importar_contas_pagar_sienge(request):
    resultado = None
    if request.method == 'POST':
        form = ImportarCredoresSiengeForm(request.POST, request.FILES)
        if form.is_valid():
            conteudo = decodificar_csv_upload(form.cleaned_data['arquivo'])
            try:
                if form.cleaned_data['tipo_relatorio'] == 'pago':
                    resultado = importar_contas_pagas_credores_csv(conteudo, empresa=request.empresa)
                else:
                    resultado = importar_contas_pagar_credores_csv(conteudo, empresa=request.empresa)
                messages.success(
                    request,
                    f'Importacao concluida: {resultado.criadas} criada(s), '
                    f'{resultado.atualizadas} atualizada(s), {resultado.ignoradas} ignorada(s).',
                )
            except DatabaseError as exc:
                messages.error(request, f'Nao foi possivel importar agora. Tente novamente em instantes. Detalhe: {exc}')
    else:
        form = ImportarCredoresSiengeForm()

    return render(
        request,
        'financeiro/importar_credores_sienge.html',
        {'form': form, 'resultado': resultado},
    )


@financeiro_required
@empresa_required
def lista_contas_pagas(request):
    contas = _base_pagar(request.empresa).filter(status=ContaPagar.STATUS_PAGO).order_by('-data_pagamento', '-id')
    page_obj = _paginar(request, contas)
    return render(
        request,
        'financeiro/lista_contas_pagar.html',
        {
            'contas': page_obj,
            'page_obj': page_obj,
            'titulo': 'Contas Pagas',
            'descricao': 'Historico das despesas ja baixadas.',
            'mostrar_acoes_massa': False,
        },
    )


@financeiro_required
@empresa_required
def lista_contas_pagar_canceladas(request):
    contas = _base_pagar(request.empresa).filter(status=ContaPagar.STATUS_CANCELADO).order_by('-updated_at', '-id')
    page_obj = _paginar(request, contas)
    return render(
        request,
        'financeiro/lista_contas_pagar.html',
        {
            'contas': page_obj,
            'page_obj': page_obj,
            'titulo': 'Contas Canceladas',
            'descricao': 'Despesas canceladas e retiradas da lista principal.',
            'mostrar_acoes_massa': False,
        },
    )


@financeiro_required
@empresa_required
def lista_previsoes_financeiras(request):
    previsoes = _base_previsoes(request.empresa)
    status = request.GET.get('status')
    tipo = request.GET.get('tipo')
    if status:
        previsoes = previsoes.filter(status=status)
    if tipo:
        previsoes = previsoes.filter(tipo=tipo)
    page_obj = _paginar(request, previsoes)
    return render(
        request,
        'financeiro/lista_previsoes.html',
        {
            'previsoes': page_obj,
            'page_obj': page_obj,
            'status_atual': status or '',
            'tipo_atual': tipo or '',
        },
    )


@financeiro_required
@empresa_required
def nova_previsao_financeira(request):
    if request.method == 'POST':
        form = PrevisaoFinanceiraForm(request.POST, empresa=request.empresa)
        if form.is_valid():
            form.save()
            messages.success(request, 'Previsao financeira cadastrada com sucesso.')
            return redirect('lista_previsoes_financeiras')
    else:
        form = PrevisaoFinanceiraForm(initial={'status': PrevisaoFinanceira.STATUS_ATIVA}, empresa=request.empresa)
    return render(request, 'financeiro/form_conta.html', {'form': form, 'titulo': 'Nova Previsao Financeira'})


@financeiro_required
@empresa_required
def editar_previsao_financeira(request, previsao_id):
    previsao = get_object_or_404(PrevisaoFinanceira, id=previsao_id, empresa=request.empresa)
    if request.method == 'POST':
        form = PrevisaoFinanceiraForm(request.POST, instance=previsao, empresa=request.empresa)
        if form.is_valid():
            form.save()
            messages.success(request, 'Previsao financeira atualizada com sucesso.')
            return redirect('lista_previsoes_financeiras')
    else:
        form = PrevisaoFinanceiraForm(instance=previsao, empresa=request.empresa)
    return render(request, 'financeiro/form_conta.html', {'form': form, 'titulo': 'Editar Previsao Financeira'})


@financeiro_required
@empresa_required
@require_POST
def acao_massa_contas_pagar(request):
    ids = request.POST.getlist('contas')
    acao = request.POST.get('acao')
    data_baixa = request.POST.get('data_baixa') or timezone.localdate()
    contas = ContaPagar.objects.filter(id__in=ids, status=ContaPagar.STATUS_ABERTO, empresa=request.empresa)

    if not ids:
        messages.warning(request, 'Selecione ao menos uma conta.')
        return redirect('lista_contas_pagar')

    if acao not in {'pagar', 'cancelar'}:
        messages.warning(request, 'Escolha uma acao valida.')
        return redirect('lista_contas_pagar')

    total = 0
    for conta in contas:
        if acao == 'pagar':
            conta.status = ContaPagar.STATUS_PAGO
            conta.data_pagamento = data_baixa
        else:
            conta.status = ContaPagar.STATUS_CANCELADO
        conta.save()
        total += 1

    if acao == 'pagar':
        messages.success(request, f'{total} conta(s) marcada(s) como pagas.')
        return redirect('lista_contas_pagas')

    messages.success(request, f'{total} conta(s) cancelada(s).')
    return redirect('lista_contas_pagar_canceladas')


@financeiro_required
@empresa_required
def nova_conta_receber(request):
    if request.method == 'POST':
        form = ContaReceberForm(request.POST, empresa=request.empresa)
        if form.is_valid():
            conta = form.save()
            messages.success(request, 'Conta a receber cadastrada com sucesso.')
            return redirect('lista_contas_receber')
    else:
        initial = {}
        obra_id = request.GET.get('obra')
        if obra_id:
            initial['obra'] = obra_id
        form = ContaReceberForm(initial=initial, empresa=request.empresa)
    return render(request, 'financeiro/form_conta.html', {'form': form, 'titulo': 'Nova Conta a Receber'})


@financeiro_required
@empresa_required
def editar_conta_receber(request, conta_id):
    conta = get_object_or_404(ContaReceber, id=conta_id, empresa=request.empresa)
    if request.method == 'POST':
        form = ContaReceberForm(request.POST, instance=conta, empresa=request.empresa)
        if form.is_valid():
            form.save()
            messages.success(request, 'Conta a receber atualizada com sucesso.')
            return redirect('lista_contas_receber')
    else:
        form = ContaReceberForm(instance=conta, empresa=request.empresa)
    return render(request, 'financeiro/form_conta.html', {'form': form, 'titulo': 'Editar Conta a Receber'})


@financeiro_required
@empresa_required
def baixar_conta_receber(request, conta_id):
    conta = get_object_or_404(ContaReceber, id=conta_id, empresa=request.empresa)
    if conta.status != ContaReceber.STATUS_ABERTO:
        messages.error(request, 'Somente contas a receber em aberto podem ser recebidas.')
        return redirect('lista_contas_receber')

    if request.method == 'POST':
        form = ContaReceberBaixaForm(request.POST)
        if form.is_valid():
            try:
                baixar_conta_receber_service(conta, data_recebimento=form.cleaned_data['data_recebimento'])
                if form.cleaned_data.get('observacoes'):
                    conta.observacoes = '\n'.join(
                        value for value in [conta.observacoes, form.cleaned_data['observacoes']] if value
                    )
                    conta.save(update_fields=['observacoes', 'updated_at'])
            except ValidationError as exc:
                messages.error(request, exc.messages[0] if exc.messages else 'Nao foi possivel receber esta conta.')
                return redirect('lista_contas_receber')
            messages.success(request, 'Recebimento registrado com sucesso.')
            return redirect('lista_contas_receber')
    else:
        form = ContaReceberBaixaForm(initial={'data_recebimento': timezone.localdate()})

    return render(request, 'financeiro/form_recebimento.html', {'form': form, 'conta': conta})


@financeiro_required
@empresa_required
@require_POST
def cancelar_conta_receber(request, conta_id):
    conta = get_object_or_404(ContaReceber, id=conta_id, empresa=request.empresa)
    if conta.status != ContaReceber.STATUS_ABERTO:
        messages.error(request, 'Somente contas a receber em aberto podem ser canceladas.')
        return redirect('lista_contas_receber')
    conta.status = ContaReceber.STATUS_CANCELADO
    conta.save()
    messages.success(request, 'Conta a receber cancelada com sucesso.')
    return redirect('lista_contas_receber')


@financeiro_required
@empresa_required
def nova_conta_pagar(request):
    if request.method == 'POST':
        form = ContaPagarForm(request.POST, empresa=request.empresa)
        ordem = form.data.get('ordem_compra') or None
        formset = ItemContaPagarOrdemCompraFormSet(request.POST, ordem=ordem, empresa=request.empresa, prefix='itens_oc')
        if form.is_valid() and formset.is_valid():
            if form.cleaned_data.get('ordem_compra') and not _formset_tem_itens_oc(formset):
                form.add_error('ordem_compra', 'Informe ao menos um item da OC.')
            else:
                conta = form.save()
                formset.instance = conta
                formset.save()
                conta.recalcular_valor_por_itens_oc()
                conta.save()
                messages.success(request, 'Conta a pagar cadastrada com sucesso.')
                return redirect('lista_contas_pagar')
    else:
        initial = {}
        obra_id = request.GET.get('obra')
        if obra_id:
            initial['obra'] = obra_id
        ordem_id = request.GET.get('ordem_compra')
        if ordem_id:
            initial['ordem_compra'] = ordem_id
        form = ContaPagarForm(initial=initial, empresa=request.empresa)
        formset = ItemContaPagarOrdemCompraFormSet(ordem=ordem_id, empresa=request.empresa, prefix='itens_oc')
    return render(
        request,
        'financeiro/form_conta.html',
        {'form': form, 'item_formset': formset, 'titulo': 'Nova Conta a Pagar'},
    )


@financeiro_required
@empresa_required
def editar_conta_pagar(request, conta_id):
    conta = get_object_or_404(ContaPagar, id=conta_id, empresa=request.empresa)
    if request.method == 'POST':
        form = ContaPagarForm(request.POST, instance=conta, empresa=request.empresa)
        ordem = form.data.get('ordem_compra') or None
        formset = ItemContaPagarOrdemCompraFormSet(
            request.POST,
            instance=conta,
            ordem=ordem,
            empresa=request.empresa,
            prefix='itens_oc',
        )
        if form.is_valid() and formset.is_valid():
            if form.cleaned_data.get('ordem_compra') and not _formset_tem_itens_oc(formset):
                form.add_error('ordem_compra', 'Informe ao menos um item da OC.')
            else:
                conta = form.save()
                formset.save()
                conta.recalcular_valor_por_itens_oc()
                conta.save()
                messages.success(request, 'Conta a pagar atualizada com sucesso.')
                return redirect('lista_contas_pagar')
    else:
        form = ContaPagarForm(instance=conta, empresa=request.empresa)
        formset = ItemContaPagarOrdemCompraFormSet(
            instance=conta,
            ordem=conta.ordem_compra,
            empresa=request.empresa,
            prefix='itens_oc',
        )
    return render(
        request,
        'financeiro/form_conta.html',
        {'form': form, 'item_formset': formset, 'titulo': 'Editar Conta a Pagar'},
    )


@financeiro_required
@empresa_required
def baixar_conta_pagar(request, conta_id):
    conta = get_object_or_404(ContaPagar, id=conta_id, empresa=request.empresa)
    if conta.status != ContaPagar.STATUS_ABERTO:
        messages.error(request, 'Somente contas a pagar em aberto podem ser pagas.')
        return redirect('lista_contas_pagar')

    if request.method == 'POST':
        form = ContaPagarBaixaForm(request.POST)
        if form.is_valid():
            baixar_conta_pagar_service(
                conta,
                data_pagamento=form.cleaned_data['data_pagamento'],
                valor_pago=form.cleaned_data['valor_pago'],
            )
            if form.cleaned_data.get('observacoes'):
                conta.observacoes = '\n'.join(
                    value for value in [conta.observacoes, form.cleaned_data['observacoes']] if value
                )
                conta.save(update_fields=['observacoes', 'updated_at'])
            messages.success(request, 'Pagamento registrado com sucesso.')
            return redirect('lista_contas_pagas')
    else:
        form = ContaPagarBaixaForm(
            initial={
                'data_pagamento': timezone.localdate(),
                'valor_pago': conta.valor,
            }
        )

    return render(request, 'financeiro/form_pagamento.html', {'form': form, 'conta': conta})


@financeiro_required
@empresa_required
def lista_centros_custo(request):
    centros = CentroCusto.objects.filter(empresa=request.empresa)
    return render(request, 'financeiro/lista_centros_custo.html', {'centros': centros})


@financeiro_required
@empresa_required
def lista_fornecedores(request):
    fornecedores = Fornecedor.objects.filter(empresa=request.empresa)
    busca = request.GET.get('busca', '').strip()
    if busca:
        fornecedores = fornecedores.filter(
            Q(nome__icontains=busca)
            | Q(cpf_cnpj__icontains=busca)
            | Q(municipio__icontains=busca)
        )
    return render(request, 'financeiro/lista_fornecedores.html', {'fornecedores': fornecedores, 'busca': busca})


@financeiro_required
@empresa_required
def novo_fornecedor(request):
    if request.method == 'POST':
        form = FornecedorForm(request.POST, empresa=request.empresa)
        if form.is_valid():
            form.save()
            messages.success(request, 'Fornecedor cadastrado com sucesso.')
            return redirect('lista_fornecedores')
    else:
        form = FornecedorForm(empresa=request.empresa)
    return render(request, 'financeiro/form_conta.html', {'form': form, 'titulo': 'Novo Fornecedor'})


@financeiro_required
@empresa_required
def editar_fornecedor(request, fornecedor_id):
    fornecedor = get_object_or_404(Fornecedor, id=fornecedor_id, empresa=request.empresa)
    if request.method == 'POST':
        form = FornecedorForm(request.POST, instance=fornecedor, empresa=request.empresa)
        if form.is_valid():
            form.save()
            messages.success(request, 'Fornecedor atualizado com sucesso.')
            return redirect('lista_fornecedores')
    else:
        form = FornecedorForm(instance=fornecedor, empresa=request.empresa)
    return render(request, 'financeiro/form_conta.html', {'form': form, 'titulo': 'Editar Fornecedor'})


@financeiro_required
@empresa_required
def excluir_fornecedor(request, fornecedor_id):
    fornecedor = get_object_or_404(Fornecedor, id=fornecedor_id, empresa=request.empresa)
    detalhe = (
        'Os lancamentos e ordens ja criados nao serao apagados. '
        'Eles manterao os dados copiados do fornecedor, mas perderao o vinculo com este cadastro central.'
    )
    if request.method == 'POST':
        nome = fornecedor.nome
        fornecedor.delete()
        messages.success(request, f'Fornecedor "{nome}" excluido com sucesso.')
        return redirect('lista_fornecedores')
    return render(
        request,
        'obras/confirmar_exclusao.html',
        {
            'titulo': 'Excluir fornecedor',
            'mensagem': f'Deseja excluir o fornecedor "{fornecedor}"?',
            'detalhe': detalhe,
            'confirmar_label': 'Excluir fornecedor',
            'cancelar_href': reverse('lista_fornecedores'),
        },
    )


@financeiro_required
@empresa_required
def novo_centro_custo(request):
    if request.method == 'POST':
        form = CentroCustoForm(request.POST, empresa=request.empresa)
        if form.is_valid():
            form.save()
            messages.success(request, 'Centro de custo cadastrado com sucesso.')
            return redirect('lista_centros_custo')
    else:
        form = CentroCustoForm(empresa=request.empresa)
    return render(request, 'financeiro/form_conta.html', {'form': form, 'titulo': 'Novo Centro de Custo'})


@financeiro_required
@empresa_required
def editar_centro_custo(request, centro_id):
    centro = get_object_or_404(CentroCusto, id=centro_id, empresa=request.empresa)
    if request.method == 'POST':
        form = CentroCustoForm(request.POST, instance=centro, empresa=request.empresa)
        if form.is_valid():
            form.save()
            messages.success(request, 'Centro de custo atualizado com sucesso.')
            return redirect('lista_centros_custo')
    else:
        form = CentroCustoForm(instance=centro, empresa=request.empresa)
    return render(request, 'financeiro/form_conta.html', {'form': form, 'titulo': 'Editar Centro de Custo'})


@financeiro_required
@empresa_required
def relatorio_financeiro(request):
    data = request.GET.copy()
    if 'colunas' not in data:
        data.setlist('colunas', FinanceiroFiltroForm.COLUNAS_PADRAO)
    form, receber, pagar, previsoes = _filtrar_contas(request, data)
    ordenacao = form.cleaned_data.get('ordenacao') if form.is_valid() else 'data_asc'
    agrupamento = form.cleaned_data.get('agrupamento') if form.is_valid() else ''
    colunas = FinanceiroFiltroForm.COLUNAS_PADRAO
    eventos = _ordenar_eventos(_eventos_fluxo(receber, pagar, previsoes), ordenacao)
    grupos = _agrupar_eventos(eventos, agrupamento)
    resumo = _resumo(receber, pagar, previsoes)
    if form.is_valid():
        colunas = form.cleaned_data['colunas']
        export = request.GET.get('export')
        if export == 'excel':
            response = HttpResponse(
                _xlsx_relatorio_financeiro(eventos, colunas, request.empresa),
                content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            )
            response['Content-Disposition'] = 'attachment; filename="relatorio_financeiro.xlsx"'
            return response
        if export == 'pdf':
            response = HttpResponse(_financial_report_pdf(eventos, colunas, resumo, request.empresa), content_type='application/pdf')
            response['Content-Disposition'] = 'inline; filename="relatorio_financeiro.pdf"'
            return response
    query_params = request.GET.copy()
    query_params.pop('export', None)
    query_string = query_params.urlencode()
    contexto = {
        'filtro_form': form,
        'eventos': eventos,
        'grupos_eventos': grupos,
        'total_eventos': len(eventos),
        'ordenacao_atual': ordenacao,
        'agrupamento_atual': agrupamento,
        'colunas': colunas,
        'colunas_labels': RELATORIO_FINANCEIRO_COLUNAS,
        'query_string': query_string,
        **resumo,
    }
    return render(request, 'financeiro/relatorio.html', contexto)


@financeiro_required
@empresa_required
def relatorio_financeiro_pdf(request):
    data = request.GET.copy()
    if 'colunas' not in data:
        data.setlist('colunas', FinanceiroFiltroForm.COLUNAS_PADRAO)
    form, receber, pagar, previsoes = _filtrar_contas(request, data)
    ordenacao = form.cleaned_data.get('ordenacao') if form.is_valid() else 'data_asc'
    colunas = form.cleaned_data.get('colunas') if form.is_valid() else FinanceiroFiltroForm.COLUNAS_PADRAO
    eventos = _ordenar_eventos(_eventos_fluxo(receber, pagar, previsoes), ordenacao)
    resumo = _resumo(receber, pagar, previsoes)
    response = HttpResponse(
        _financial_report_pdf(eventos, colunas, resumo, request.empresa),
        content_type='application/pdf',
    )
    response['Content-Disposition'] = 'inline; filename="relatorio_financeiro.pdf"'
    return response
